"""Shared business logic and report builder for the Excel sales automation project.

Every entry point (CLI ``main3.py`` and the Streamlit ``app.py``) calls the
functions in this module, so the generated report is always identical.

Public API
----------
BASE_DIR / INPUT_FILE / OUTPUT_FILE   project paths (resolved, not CWD based)
DataValidationError                   raised for unreadable / invalid input
load_sales_data(source)               read + validate + clean an Excel file
clean_sales_data(df)                  validate + clean a DataFrame
summarize(df)                         compute every metric used in the report
write_report(df, summary, target)     render the full styled workbook
build_report(source, target=None)     load_sales_data + summarize + write_report
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import IO, Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Paths — always resolved from this file, never from the current directory.
# --------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
INPUT_FILE = INPUT_DIR / "sales.xlsx"
OUTPUT_FILE = OUTPUT_DIR / "premium_automated_report.xlsx"

REQUIRED_COLUMNS = ["Product", "Region", "Quantity", "Price"]
DEFAULT_CURRENCY = "$"

# Cells below this limit get full border styling; bigger sheets keep only the
# header styling so the workbook stays fast and small.
STYLED_ROWS_LIMIT = 5000


class DataValidationError(ValueError):
    """The input file cannot be turned into a report."""


Target = Union[str, Path, IO[bytes]]


# --------------------------------------------------------------------------
# Loading & cleaning
# --------------------------------------------------------------------------
def load_sales_data(source: Union[str, Path, IO[bytes]]) -> pd.DataFrame:
    """Read an Excel file (path or file-like) and return cleaned sales data."""
    try:
        raw = pd.read_excel(source)
    except DataValidationError:
        raise
    except Exception as exc:  # noqa: BLE001 - surfaced to the UI as one message
        raise DataValidationError(f"Could not read the Excel file: {exc}") from exc
    return clean_sales_data(raw)


def clean_sales_data(df: pd.DataFrame) -> pd.DataFrame:
    """Validate columns, normalise types and drop unusable rows."""
    if df is None or df.empty:
        raise DataValidationError("The file contains no rows.")

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise DataValidationError(
            f"Missing required columns: {missing}. Expected {REQUIRED_COLUMNS}."
        )

    out = df.copy()

    # "string" dtype keeps missing values as <NA> instead of the literal "nan",
    # so the dropna() below actually removes rows with a blank Product/Region.
    for col in ("Product", "Region"):
        out[col] = out[col].astype("string").str.strip()
        out[col] = out[col].replace("", pd.NA)

    for col in ("Quantity", "Price"):
        out[col] = pd.to_numeric(out[col], errors="coerce")

    before = len(out)
    out = out.dropna(subset=REQUIRED_COLUMNS)
    out = out[(out["Quantity"] > 0) & (out["Price"] > 0)]
    out = out.reset_index(drop=True)

    if out.empty:
        raise DataValidationError(
            f"No valid rows left after cleaning (discarded {before} of {before} rows). "
            "Quantity and Price must be positive numbers."
        )

    out["Revenue"] = out["Quantity"] * out["Price"]
    return out


# --------------------------------------------------------------------------
# Metrics
# --------------------------------------------------------------------------
@dataclass(frozen=True)
class SalesSummary:
    total_revenue: float
    total_quantity: float
    average_unit_price: float
    weighted_price: float
    transactions: int
    best_product: str
    best_product_revenue: float
    best_region: str
    best_region_revenue: float
    product_sales: pd.Series  # revenue per product, descending
    region_sales: pd.Series   # revenue per region, descending

    def top_products(self, n: int = 5) -> pd.Series:
        return self.product_sales.head(n)

    def top_regions(self, n: int = 5) -> pd.Series:
        return self.region_sales.head(n)


def summarize(df: pd.DataFrame) -> SalesSummary:
    """Compute every number shown in the dashboard."""
    if df.empty:
        raise DataValidationError("Cannot summarise an empty dataset.")

    product_sales = df.groupby("Product", dropna=False)["Revenue"].sum().sort_values(ascending=False)
    region_sales = df.groupby("Region", dropna=False)["Revenue"].sum().sort_values(ascending=False)
    total_revenue = float(df["Revenue"].sum())

    return SalesSummary(
        total_revenue=total_revenue,
        total_quantity=float(df["Quantity"].sum()),
        average_unit_price=float(df["Price"].mean()),
        weighted_price=total_revenue / float(df["Quantity"].sum()),
        transactions=int(len(df)),
        best_product=str(product_sales.index[0]),
        best_product_revenue=float(product_sales.iloc[0]),
        best_region=str(region_sales.index[0]),
        best_region_revenue=float(region_sales.iloc[0]),
        product_sales=product_sales,
        region_sales=region_sales,
    )


def money(value: float, currency: str = DEFAULT_CURRENCY) -> str:
    return f"{currency}{value:,.2f}"


# --------------------------------------------------------------------------
# Report builder
# --------------------------------------------------------------------------
# Palette
NAVY = "17324D"
BLUE = "2F6690"
LIGHT_BLUE = "EAF2F8"
WHITE = "FFFFFF"
DARK = "1F2937"
BORDER_COLOR = "D9E1E8"
GREEN = "1F7A5C"
GREEN_LIGHT = "EAF7F2"
MUTED = "64748B"


def _money_format(currency: str) -> str:
    # '"$"#,##0.00' keeps the symbol literal so any currency string works.
    return f'"{currency}"#,##0.00'


def build_report(
    source: Union[str, Path, IO[bytes]],
    target: Target | None = None,
    *,
    currency: str = DEFAULT_CURRENCY,
) -> tuple[SalesSummary, bytes]:
    """Load ``source``, build the styled workbook and write it to ``target``."""
    df = load_sales_data(source)
    summary = summarize(df)
    data = write_report(df, summary, target, currency=currency)
    return summary, data


def write_report(
    df: pd.DataFrame,
    summary: SalesSummary,
    target: Target | None = None,
    *,
    currency: str = DEFAULT_CURRENCY,
) -> bytes:
    """Render the full workbook in memory, write it to ``target``, return bytes."""
    money_fmt = _money_format(currency)
    excel_bytes = _render_workbook(df, summary, currency, money_fmt)

    if target is None:
        return excel_bytes
    if isinstance(target, (str, Path)):
        path = Path(target)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(excel_bytes)
    else:
        target.write(excel_bytes)
    return excel_bytes


def _render_workbook(
    df: pd.DataFrame,
    summary: SalesSummary,
    currency: str,
    money_fmt: str,
) -> bytes:
    # ---- data sheets -----------------------------------------------------
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sales Data", index=False)
        (
            summary.product_sales
            .rename("Revenue")
            .rename_axis("Product")  # otherwise column A header is blank
            .to_excel(writer, sheet_name="By Product")
        )
        (
            summary.region_sales
            .rename("Revenue")
            .rename_axis("Region")
            .to_excel(writer, sheet_name="By Region")
        )
        pd.DataFrame(
            {
                "Metric": [
                    "Total Revenue",
                    "Total Quantity",
                    "Average Unit Price",
                    "Weighted Price",
                    "Total Transactions",
                    "Best Product",
                    "Best Region",
                ],
                "Value": [
                    summary.total_revenue,
                    summary.total_quantity,
                    summary.average_unit_price,
                    summary.weighted_price,
                    summary.transactions,
                    summary.best_product,
                    summary.best_region,
                ],
            }
        ).to_excel(writer, sheet_name="Business Insights", index=False)

    buffer.seek(0)
    wb = load_workbook(buffer)

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    _build_dashboard(wb, summary, currency, money_fmt, border)
    _style_data_sheets(wb, money_fmt, border)
    _autosize(wb)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _build_dashboard(wb, summary: SalesSummary, currency: str, money_fmt: str, border) -> None:
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False

    # white canvas first — must happen before merging cells
    for row in ws.iter_rows(min_row=1, max_row=30, min_col=1, max_col=14):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=WHITE)

    ws.merge_cells("A1:N2")
    ws["A1"] = "AI SALES AUTOMATION"
    ws["A1"].font = Font(size=24, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:N3")
    ws["A3"] = "Professional Sales Analytics Dashboard"
    ws["A3"].font = Font(size=11, italic=True, color=MUTED)
    ws["A3"].alignment = Alignment(horizontal="center")

    cards = [
        ("A5:C5", "A6:C8", "TOTAL REVENUE", summary.total_revenue, money_fmt),
        ("D5:F5", "D6:F8", "TOTAL QUANTITY", summary.total_quantity, "#,##0"),
        ("G5:I5", "G6:I8", "AVG UNIT PRICE", summary.average_unit_price, money_fmt),
        ("J5:L5", "J6:L8", "TRANSACTIONS", summary.transactions, "#,##0"),
    ]
    for title_range, value_range, title, value, numfmt in cards:
        _card(ws, title_range, value_range, title, value, numfmt, border)

    highlights = [
        (
            "A10:F10",
            "A11:F13",
            "TOP PRODUCT",
            f"{summary.best_product}\n{money(summary.best_product_revenue, currency)}",
        ),
        (
            "G10:L10",
            "G11:L13",
            "TOP REGION",
            f"{summary.best_region}\n{money(summary.best_region_revenue, currency)}",
        ),
    ]
    for title_range, value_range, title, value in highlights:
        _highlight(ws, title_range, value_range, title, value, border)

    # ---- charts ----------------------------------------------------------
    product_ws = wb["By Product"]
    region_ws = wb["By Region"]

    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.title = "Revenue by Product"
    bar.y_axis.title = "Product"
    bar.x_axis.title = "Revenue"
    bar.add_data(
        Reference(product_ws, min_col=2, min_row=1, max_row=len(summary.product_sales) + 1),
        titles_from_data=True,
    )
    bar.set_categories(
        Reference(product_ws, min_col=1, min_row=2, max_row=len(summary.product_sales) + 1)
    )
    bar.height = 8
    bar.width = 14
    ws.add_chart(bar, "A16")

    pie = PieChart()
    pie.title = "Revenue by Region"
    pie.add_data(
        Reference(region_ws, min_col=2, min_row=1, max_row=len(summary.region_sales) + 1),
        titles_from_data=True,
    )
    pie.set_categories(
        Reference(region_ws, min_col=1, min_row=2, max_row=len(summary.region_sales) + 1)
    )
    pie.height = 8
    pie.width = 12
    ws.add_chart(pie, "H16")

    # ---- layout ----------------------------------------------------------
    for col in "ABCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 13
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30
    for r in (6, 7, 8):
        ws.row_dimensions[r].height = 25


def _card(ws, title_range: str, value_range: str, title: str, value, numfmt: str, border) -> None:
    ws.merge_cells(title_range)
    ws.merge_cells(value_range)
    title_cell = ws[title_range.split(":")[0]]
    value_cell = ws[value_range.split(":")[0]]
    title_cell.value = title
    value_cell.value = value

    title_cell.fill = PatternFill("solid", fgColor=BLUE)
    title_cell.font = Font(size=10, bold=True, color=WHITE)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    value_cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    value_cell.font = Font(size=19, bold=True, color=DARK)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.number_format = numfmt

    _frame(
        ws,
        title_cell.row,
        ws[value_range.split(":")[1]].row,
        title_cell.column,
        ws[title_range.split(":")[1]].column,
        border,
    )


def _highlight(ws, title_range: str, value_range: str, title: str, value: str, border) -> None:
    ws.merge_cells(title_range)
    ws.merge_cells(value_range)
    title_cell = ws[title_range.split(":")[0]]
    value_cell = ws[value_range.split(":")[0]]
    title_cell.value = title
    value_cell.value = value

    title_cell.fill = PatternFill("solid", fgColor=GREEN)
    title_cell.font = Font(size=10, bold=True, color=WHITE)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    value_cell.fill = PatternFill("solid", fgColor=GREEN_LIGHT)
    value_cell.font = Font(size=18, bold=True, color=DARK)
    value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    _frame(
        ws,
        title_cell.row,
        ws[value_range.split(":")[1]].row,
        title_cell.column,
        ws[title_range.split(":")[1]].column,
        border,
    )


def _frame(ws, min_row: int, max_row: int, min_col: int, max_col: int, border) -> None:
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(r, c).border = border


def _style_data_sheets(wb, money_fmt: str, border) -> None:
    for sheet_name in ["Business Insights", "Sales Data", "By Product", "By Region"]:
        sh = wb[sheet_name]
        sh.sheet_view.showGridLines = False
        sh.freeze_panes = "A2"

        for cell in sh[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center")

        if sh.max_row <= STYLED_ROWS_LIMIT:
            for row in sh.iter_rows():
                for cell in row:
                    cell.border = border

    # money columns on the raw data sheet
    sales_ws = wb["Sales Data"]
    headers = {cell.value: cell.column for cell in sales_ws[1]}
    for col in ("Price", "Revenue"):
        if col in headers:
            for r in range(2, sales_ws.max_row + 1):
                sales_ws.cell(r, headers[col]).number_format = money_fmt

    for shname in ("By Product", "By Region"):
        sh = wb[shname]
        for r in range(2, sh.max_row + 1):
            sh.cell(r, 2).number_format = money_fmt

    _style_insights(wb["Business Insights"], money_fmt)


def _style_insights(ins, money_fmt: str) -> None:
    """Format the metric values by their row instead of hard-coded cells."""
    # rows: 2 Total Revenue, 3 Total Quantity, 4 Average Unit Price,
    #       5 Weighted Price, 6 Transactions, 7 Best Product, 8 Best Region
    formats = {
        2: money_fmt,
        3: "#,##0",
        4: money_fmt,
        5: money_fmt,
        6: "#,##0",
    }
    for row, numfmt in formats.items():
        ins.cell(row, 2).number_format = numfmt
    ins.cell(2, 2).font = Font(bold=True, color=GREEN)
    ins.cell(7, 2).font = Font(bold=True, color=DARK)
    ins.cell(8, 2).font = Font(bold=True, color=DARK)


def _autosize(wb) -> None:
    for sh in wb.worksheets:
        if sh.max_row > STYLED_ROWS_LIMIT:
            # only measure the header row on huge sheets
            cells = list(sh[1])
        else:
            cells = [c for col in sh.columns for c in col]
        for cell in cells:
            if cell.value is None:
                continue
            letter = get_column_letter(cell.column)
            length = len(str(cell.value))
            current = sh.column_dimensions[letter].width or 0
            sh.column_dimensions[letter].width = min(max(max(current, length + 3), 12), 40)
